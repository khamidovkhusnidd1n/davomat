import docx

file_path = r"C:\Users\Salohiddin Markaz\Desktop\davomat\Davomat_Mustaqil_Ish_Raqamli_Ozgarish.docx"
doc = docx.Document(file_path)

xulosa_idx = -1
for i, p in enumerate(doc.paragraphs):
    if "XULOSA" in p.text and len(p.text) < 20:
        xulosa_idx = i
        break

if xulosa_idx != -1:
    xulosa_p = doc.paragraphs[xulosa_idx]
    
    def add_heading(text, level, p_ref):
        new_p = p_ref.insert_paragraph_before(text)
        try:
            new_p.style = doc.styles[f'Heading {level}']
        except:
            pass
        return new_p
        
    def add_para(text, p_ref):
        return p_ref.insert_paragraph_before(text)

    add_para('\n', xulosa_p)
    add_heading("4-BOB. STARTAP G'OYASINI SHAKLLANTIRISH: 'IDEA TO MARKET' MODELI", 1, xulosa_p)
    add_para("“Idea to Market” (G'oyadan bozorgacha) modeli asosida ushbu davomat tizimini tijoratlashtirish va muvaffaqiyatli startapga aylantirish bosqichlari ko'rib chiqiladi.", xulosa_p)
    
    add_heading("4.1. Canvas modeli tuzish (Business Model Canvas)", 2, xulosa_p)
    add_para("Loyihaning biznes modeli quyidagi 9 ta asosiy blokdan iborat:\n"
             "1. Mijozlar segmenti (Customer Segments): O'quv markazlari, xususiy maktablar va kichik/o'rta biznes korxonalari.\n"
             "2. Qiymat taklifi (Value Propositions): Davomat jarayonini 90% ga tezlashtirish, qog'ozsiz ish yuritish (Zero Paper), real vaqt rejimida monitoring va avtomatik aniq hisobotlar.\n"
             "3. Kanallar (Channels): B2B to'g'ridan-to'g'ri sotuvlar (Direct Sales), ijtimoiy tarmoqlar va ta'lim texnologiyalari ko'rgazmalari.\n"
             "4. Mijozlar bilan munosabat (Customer Relationships): 24/7 onlayn texnik yordam, platformadan bepul o'qitish va doimiy tizim yangilanishlari.\n"
             "5. Daromad oqimi (Revenue Streams): SaaS modeli asosida oylik/yillik obuna to'lovlari. Ta'lim muassasasidagi o'quvchilar soniga qarab moslashuvchan tariflar.\n"
             "6. Asosiy resurslar (Key Resources): Malakali IT mutaxassislar, xavfsiz bulutli serverlar (Supabase) va kuchli savdo menejerlari jamoasi.\n"
             "7. Asosiy harakatlar (Key Activities): Dasturiy ta'minotni ishlab chiqish, server barqarorligini ta'minlash, hamda marketing kampaniyalari.\n"
             "8. Asosiy hamkorlar (Key Partners): Ta'lim vazirligi tizimidagi boshqarmalar, mahalliy to'lov tizimlari (Payme, Click) va Cloud xizmatlari provayderlari.\n"
             "9. Xarajatlar tuzilmasi (Cost Structure): Server ijarasi, jamoa oylik maoshlari va marketing xarajatlari.\n", xulosa_p)

    add_heading("4.2. Minimal Viable Product (MVP) konsepsiyasi", 2, xulosa_p)
    add_para("Har qanday startap bozorda o'z o'rnini topishi uchun dastlab Minimal Viable Product (Eng zarur funksiyalarga ega minimal mahsulot) ishlab chiqilishi shart. Davomat tizimining MVP versiyasi quyidagilarni o'z ichiga oladi:\n"
             "- Telegram Bot orqali faqatgina eng zarur bo'lgan 'Keldi' yoki 'Kelmadi' maqomini kiritish imkoniyati.\n"
             "- Admin panelda faqat bugungi kun statistikasi va oddiy Excel formatda eksport qilish funksiyasi.\n"
             "- Supabase orqali ruxsatlarni (login/parol) oddiy boshqarish mexanizmi.\n"
             "MVP yaratishdan maqsad: Dastlabki mijozlarga tizimni bepul sinov muddati bilan taqdim etish orqali bozor reaksiyasini tezkor o'lchash. Foydalanuvchilardan kelgan fikr-mulohazalar asosida keyingi murakkab modullar asta-sekin qo'shib boriladi.\n", xulosa_p)
             
    add_heading("5-BOB. YANGI MAHSULOT ISHLAB CHIQISH LOYIHASINI YARATISH", 1, xulosa_p)
    add_heading("5.1. Bozor problemi va yechimini aniqlash", 2, xulosa_p)
    add_para("Bozor problemi (Problem): O'zbekistondagi aksariyat nodavlat ta'lim muassasalari davomat hisobini hamon qog'oz jurnallar yoki umumiy Excel jadvallar orqali yuritadi. Bu quyidagi og'riqli muammolarni keltirib chiqaradi:\n"
             "- Ma'lumotlarning ishonchliligi pastligi va o'zgarib qolish xavfi.\n"
             "- O'qituvchilarning asosiy e'tibori ta'limga emas, balki qog'oz to'ldirishga chalg'ishi.\n"
             "- Rahbariyatning real vaqtda muassasa holatini tahlil qila olmasligi.\n"
             "Yechim (Solution): Telegram messenjeriga to'liq integratsiya qilingan bulutli Davomat platformasi. O'qituvchilar hech qanday yangi og'ir mobil ilovalarni yuklab olishmaydi, balki o'zlari har kuni ishlatadigan Telegram orqali 3 tugmani bosish orqali davomatni kiritadilar. Barcha tahliliy jarayonlar inson omilisiz serverda bajariladi.\n", xulosa_p)

    add_heading("5.2. Mahsulot konsepsiyasi va prototip sxemasi", 2, xulosa_p)
    add_para("Mahsulot konsepsiyasi: Loyiha 'Bir marta bosish orqali davomat va avtomatik hisobotlar' konsepsiyasiga asoslangan. Tizim an'anaviy sotilmaydi, balki bulutli xizmat (SaaS) sifatida har oylik to'lov evaziga ijaraga beriladi.\n"
             "Prototip sxemasi va Tizim arxitekturasi:\n"
             "Mahsulotning prototipi 4 ta asosiy qatlamdan iborat tarzda sxemalashtirilgan:\n"
             "1. Foydalanuvchi qatlami (User Tier): O'qituvchi smartfonida Telegram Botga kirib, guruhni tanlaydi va o'quvchilarni holatini belgilaydi.\n"
             "2. Mantiqiy qatlam (Logic Tier): Telegraf API (Node.js) ma'lumotlarni tezkor qabul qilib, tahlilga tayyorlaydi.\n"
             "3. Ma'lumotlar qatlami (Data Tier): Supabase (PostgreSQL) bazasi kelgan ma'lumotni xavfsiz va shifrlangan holda saqlaydi.\n"
             "4. Tahlil va Boshqaruv qatlami (Analytics Tier): Next.js orqali ishlangan Admin Panel ma'lumotlarni jadvallar va vizual grafiklar ko'rinishida rahbariyatga namoyish etadi.\n", xulosa_p)
    add_para('\n', xulosa_p)
    
    doc.save(file_path)
    print("Muvaffaqiyatli! Matn XULOSA dan oldin kiritildi.")
else:
    print("XULOSA topilmadi. Iltimos tekshiring.")
